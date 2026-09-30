import json
import os
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

try:
    from .formula_engine import FormulaError, calculate_formula
except ImportError:
    from formula_engine import FormulaError, calculate_formula


DEFAULT_SHEETS = {
    "Attendance": [
        {"name": "Name", "type": "text"},
        {"name": "Date", "type": "date"},
        {"name": "Status", "type": "dropdown", "options": ["Present", "Absent", "Leave"]},
        {"name": "Notes", "type": "text"},
    ],
    "Payment Record": [
        {"name": "Customer", "type": "text"},
        {"name": "Date", "type": "date"},
        {"name": "Bill", "type": "number"},
        {"name": "Paid", "type": "number"},
        {"name": "Balance", "type": "formula", "formula": "Bill - Paid"},
    ],
    "Company Details": [
        {"name": "Company", "type": "text"},
        {"name": "Contact", "type": "text"},
        {"name": "Phone", "type": "text"},
        {"name": "Email", "type": "text"},
    ],
    "Follow Up": [
        {"name": "Customer", "type": "text"},
        {"name": "Due Date", "type": "date"},
        {"name": "Status", "type": "dropdown", "options": ["Pending", "Done", "Reschedule"]},
        {"name": "Notes", "type": "text"},
    ],
}
ROW_ID_HEADER = "__ledgerly_record_id"


class ExcelManager:
    def __init__(self, workbook_path: Path, config_path: Path):
        self.workbook_path = workbook_path
        self.config_path = config_path
        self.s3_bucket = os.getenv("S3_BUCKET")
        self.s3_workbook_key = os.getenv("S3_WORKBOOK_KEY", "business-manager/Business.xlsx")
        self.s3_config_key = os.getenv("S3_CONFIG_KEY", "business-manager/config.json")
        self.s3_attachments_prefix = os.getenv("S3_ATTACHMENTS_PREFIX", "business-manager/attachments")
        self.s3_client = None
        if self.s3_bucket:
            try:
                import boto3
                from botocore.config import Config
            except ImportError as error:
                raise RuntimeError("Install boto3 to use S3-backed workbook storage") from error
            self.s3_client = boto3.client(
                "s3",
                region_name=os.getenv("S3_REGION") or None,
                endpoint_url=os.getenv("S3_ENDPOINT_URL") or None,
                config=Config(s3={"addressing_style": "path"}) if os.getenv("S3_ENDPOINT_URL") else None,
            )
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        self.config_path.parent.mkdir(parents=True, exist_ok=True)
        self.attachment_root = self.workbook_path.parent / "attachments"
        self.attachment_root.mkdir(parents=True, exist_ok=True)
        self._refresh_storage_files()
        self.config = self._load_config()
        if not self.workbook_path.exists():
            workbook = Workbook()
            workbook.remove(workbook.active)
            for name, fields in DEFAULT_SHEETS.items():
                workbook.create_sheet(name).append([field["name"] for field in fields])
                self.config["sheets"][name] = deepcopy(fields)
            self._save_workbook(workbook)
            self._save_config()
        self._sync_config_with_workbook()
        self._ensure_record_ids()

    def _load_config(self) -> dict[str, Any]:
        if self.config_path.exists():
            try:
                config = json.loads(self.config_path.read_text(encoding="utf-8"))
                if isinstance(config.get("sheets"), dict):
                    return config
            except (json.JSONDecodeError, OSError):
                pass
        return {"sheets": {}, "styles": {}}

    def _save_config(self) -> None:
        self.config_path.write_text(json.dumps(self.config, indent=2), encoding="utf-8")
        self._upload_storage_file(self.config_path, self.s3_config_key)

    def _save_workbook(self, workbook) -> None:
        workbook.save(self.workbook_path)
        self._upload_storage_file(self.workbook_path, self.s3_workbook_key)

    def _upload_storage_file(self, path: Path, key: str) -> None:
        if self.s3_client and self.s3_bucket:
            self.s3_client.upload_file(str(path), self.s3_bucket, key)

    def _refresh_storage_files(self) -> None:
        if not self.s3_client or not self.s3_bucket:
            return
        for key, path in ((self.s3_workbook_key, self.workbook_path), (self.s3_config_key, self.config_path)):
            try:
                self.s3_client.download_file(self.s3_bucket, key, str(path))
            except Exception as error:
                code = getattr(error, "response", {}).get("Error", {}).get("Code")
                if code not in {"404", "NoSuchKey", "NotFound"}:
                    raise

    def refresh_from_storage(self) -> None:
        self._refresh_storage_files()
        if self.config_path.exists():
            self.config = self._load_config()
            if self.workbook_path.exists():
                self._sync_config_with_workbook()
                self._ensure_record_ids()

    def _load_workbook(self):
        return load_workbook(self.workbook_path)

    def _sync_config_with_workbook(self) -> None:
        workbook = self._load_workbook()
        changed = False
        for sheet in workbook.worksheets:
            fields = self.config["sheets"].get(sheet.title)
            names = [cell.value for cell in sheet[1] if cell.value is not None and cell.value != ROW_ID_HEADER]
            if not fields:
                fields = [{"name": str(name), "type": "text"} for name in names]
                self.config["sheets"][sheet.title] = fields
                changed = True
            elif [field["name"] for field in fields] != names:
                sheet.delete_rows(1, 1)
                sheet.insert_rows(1)
                for index, field in enumerate(fields, start=1):
                    sheet.cell(1, index, field["name"])
                changed = True
        if changed:
            self._save_workbook(workbook)
        if changed or not self.config_path.exists():
            self._save_config()

    def _ensure_record_ids(self, workbook=None) -> None:
        workbook = workbook or self._load_workbook()
        changed = False
        for sheet in workbook.worksheets:
            fields = self.config["sheets"].get(sheet.title, [])
            if any(field["name"] == ROW_ID_HEADER for field in fields):
                raise ValueError(f"'{ROW_ID_HEADER}' is reserved for internal record identifiers")
            id_column = len(fields) + 1
            id_cell = sheet.cell(1, id_column)
            if id_cell.value != ROW_ID_HEADER:
                id_cell.value = ROW_ID_HEADER
                changed = True
            sheet.column_dimensions[id_cell.column_letter].hidden = True
            seen = set()
            for row_index in range(2, sheet.max_row + 1):
                if not any(sheet.cell(row_index, index).value is not None for index in range(1, len(fields) + 1)):
                    continue
                cell = sheet.cell(row_index, id_column)
                record_id = str(cell.value or "")
                if not record_id or record_id in seen:
                    record_id = str(uuid4())
                    cell.value = record_id
                    changed = True
                seen.add(record_id)
        if changed:
            self._save_workbook(workbook)

    def _fields(self, sheet_name: str) -> list[dict[str, Any]]:
        if sheet_name not in self.config["sheets"]:
            raise KeyError(f"Unknown list: {sheet_name}")
        return self.config["sheets"][sheet_name]

    def _read_rows(self, sheet_name: str) -> list[dict[str, Any]]:
        workbook = self._load_workbook()
        sheet = workbook[sheet_name]
        names = [field["name"] for field in self._fields(sheet_name)]
        rows = []
        id_column = len(names)
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not any(values[index] is not None for index in range(min(len(names), len(values)))):
                continue
            record = {}
            for index, name in enumerate(names):
                value = values[index] if index < len(values) else None
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()[:10]
                record[name] = value
            record["_id"] = str(values[id_column]) if id_column < len(values) and values[id_column] else ""
            rows.append(record)
        return rows

    @staticmethod
    def _summary(rows: list[dict[str, Any]], fields: list[dict[str, Any]]) -> dict[str, Any]:
        sums = {}
        for field in fields:
            if field["type"] == "number":
                values = [row.get(field["name"]) for row in rows]
                sums[field["name"]] = sum(value for value in values if isinstance(value, (int, float)))
        return {"count": len(rows), "sums": sums}

    def get_lists(self) -> list[dict[str, Any]]:
        result = []
        for name, fields in self.config["sheets"].items():
            rows = self._read_rows(name)
            result.append({"name": name, "fields": fields, **self._summary(rows, fields)})
        return result

    def get_data(self, sheet_name: str) -> dict[str, Any]:
        fields = self._field_views(self._fields(sheet_name))
        rows = self._read_rows(sheet_name)
        return {"name": sheet_name, "fields": fields, "rows": rows, **self._summary(rows, fields)}

    def _field_views(self, fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
        result = deepcopy(fields)
        for field in result:
            if field["type"] == "link":
                field["options"] = self._link_options(field)
        return result

    def _link_options(self, field: dict[str, Any]) -> list[dict[str, str]]:
        target_sheet = field.get("target_sheet")
        target_field = field.get("target_field")
        if target_sheet not in self.config["sheets"]:
            return []
        if target_field not in {item["name"] for item in self._fields(target_sheet)}:
            return []
        return [
            {"id": str(row["_id"]), "label": str(row.get(target_field) or "Untitled record")}
            for row in self._read_rows(target_sheet)
        ]

    def import_workbook(self, content: bytes, options: dict[str, Any] | None = None) -> list[dict[str, Any]]:
        old_sheets = deepcopy(self.config.get("sheets", {}))
        old_styles = deepcopy(self.config.get("styles", {}))
        old_attachments = deepcopy(self.config.get("attachments", {}))
        workbook, sheets, _ = self._inspect_workbook(content, options)
        styles = {name: {} for name in sheets}
        for sheet_name, fields in sheets.items():
            sheet = workbook[sheet_name]
            has_ids = sheet.cell(1, len(fields) + 1).value == ROW_ID_HEADER
            if not has_ids:
                continue
            previous = {field["name"]: field for field in old_sheets.get(sheet_name, [])}
            for index, field in enumerate(fields):
                prior = previous.get(field["name"])
                if prior:
                    fields[index] = deepcopy(prior)
                    if fields[index].get("type") == "link" and fields[index].get("target_sheet") not in sheets:
                        fields[index] = {"name": field["name"], "type": "text"}
                styles[sheet_name][field["name"]] = old_styles.get(sheet_name, {}).get(field["name"], {})
        self.config = {"sheets": sheets, "styles": styles}
        self._ensure_record_ids(workbook)
        imported_ids = {
            str(sheet.cell(row_index, len(sheets[sheet.title]) + 1).value)
            for sheet in workbook.worksheets
            for row_index in range(2, sheet.max_row + 1)
            if sheet.cell(row_index, len(sheets[sheet.title]) + 1).value
        }
        self._save_workbook(workbook)
        retained_attachments = {}
        for record_id, items in old_attachments.items():
            if record_id in imported_ids:
                retained_attachments[record_id] = items
            else:
                for item in items:
                    self._delete_attachment_blob(item["id"])
        if retained_attachments:
            self.config["attachments"] = retained_attachments
        self._save_config()
        return self.get_lists()

    def preview_workbook(self, content: bytes) -> dict[str, Any]:
        _, _, preview = self._inspect_workbook(content)
        return preview

    def _inspect_workbook(self, content: bytes, options: dict[str, Any] | None = None):
        try:
            workbook = load_workbook(BytesIO(content), data_only=False)
        except Exception as error:
            raise ValueError("Upload a valid .xlsx Excel workbook") from error
        if not workbook.worksheets:
            raise ValueError("The workbook must contain at least one worksheet")

        sheets = {}
        previews = []
        requested_sheets = options.get("sheets") if options else None
        if requested_sheets is not None:
            if not isinstance(requested_sheets, list) or not requested_sheets or not all(isinstance(name, str) for name in requested_sheets):
                raise ValueError("Select at least one worksheet to import")
            if set(requested_sheets) - set(workbook.sheetnames):
                raise ValueError("A selected worksheet is not in the uploaded workbook")
        column_mapping = options.get("mapping", {}) if options else {}
        if not isinstance(column_mapping, dict):
            raise ValueError("Column mappings must be an object")
        for sheet_name, mapping in column_mapping.items():
            if sheet_name not in workbook.sheetnames or not isinstance(mapping, dict):
                raise ValueError("Each column mapping must refer to an uploaded worksheet")
        for sheet in workbook.worksheets:
            if requested_sheets is not None and sheet.title not in requested_sheets:
                continue
            headers = []
            record_id_column = None
            for index, cell in enumerate(sheet[1], start=1):
                value = str(cell.value).strip() if cell.value is not None else ""
                original = value or f"Column {index}"
                if original == ROW_ID_HEADER:
                    if record_id_column is not None:
                        raise ValueError("The workbook contains duplicate internal record ID columns")
                    record_id_column = index
                    continue
                sheet_mapping = column_mapping.get(sheet.title, {})
                if not isinstance(sheet_mapping, dict):
                    raise ValueError(f"Column mappings for '{sheet.title}' must be an object")
                header = str(sheet_mapping.get(original, original)).strip()
                if not header:
                    raise ValueError(f"Column names in worksheet '{sheet.title}' cannot be empty")
                headers.append(header)
                if header != original:
                    cell.value = header
            if record_id_column is not None and record_id_column != sheet.max_column:
                raise ValueError("The internal record ID column must remain the final worksheet column")
            while headers and headers[-1].startswith("Column ") and sheet.max_column < len(headers):
                headers.pop()
            if not headers:
                raise ValueError(f"Worksheet '{sheet.title}' must have a header row")
            if len({header.casefold() for header in headers}) != len(headers):
                raise ValueError(f"Worksheet '{sheet.title}' has duplicate column names")
            if ROW_ID_HEADER in headers:
                raise ValueError(f"'{ROW_ID_HEADER}' is reserved for internal record identifiers")
            fields = []
            preview_rows = []
            row_count = 0
            for column_index, header in enumerate(headers, start=1):
                values = [sheet.cell(row, column_index).value for row in range(2, sheet.max_row + 1)]
                fields.append({"name": header, "type": self._infer_field_type(values)})
            sheets[sheet.title] = fields
            for values in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
                if not any(value is not None for value in values):
                    continue
                row_count += 1
                if len(preview_rows) < 8:
                    preview_rows.append(dict(zip(headers, values)))
            previews.append({"name": sheet.title, "fields": fields, "count": row_count, "rows": preview_rows})
        if requested_sheets is not None:
            for sheet in list(workbook.worksheets):
                if sheet.title not in requested_sheets:
                    workbook.remove(sheet)
        return workbook, sheets, {"worksheets": previews}

    @staticmethod
    def _infer_field_type(values: list[Any]) -> str:
        populated = [value for value in values if value is not None and value != ""]
        if populated and all(isinstance(value, (datetime, date)) for value in populated):
            return "date"
        if populated and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in populated):
            return "number"
        return "text"

    def _save_rows(self, sheet_name: str, rows: list[dict[str, Any]]) -> None:
        fields = self._fields(sheet_name)
        workbook = self._load_workbook()
        sheet = workbook[sheet_name]
        for index in range(1, sheet.max_column + 1):
            sheet.cell(1, index).value = None
        for index, field in enumerate(fields, start=1):
            sheet.cell(1, index, field["name"])
        id_column = len(fields) + 1
        sheet.cell(1, id_column, ROW_ID_HEADER)
        sheet.column_dimensions[sheet.cell(1, id_column).column_letter].hidden = True
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for row in rows:
            row_values = [row.get(field["name"]) for field in fields]
            row_values.append(row.get("_id") or str(uuid4()))
            sheet.append(row_values)
        self._apply_styles(workbook, sheet_name)
        self._save_workbook(workbook)

    def _calculate_formulas(self, sheet_name: str, rows: list[dict[str, Any]]) -> None:
        fields = self._fields(sheet_name)
        field_names = [field["name"] for field in fields]
        for field in fields:
            if field["type"] != "formula":
                continue
            for row in rows:
                try:
                    row[field["name"]] = calculate_formula(field.get("formula", ""), row, rows, field_names)
                except FormulaError as error:
                    raise ValueError(f"Formula for {field['name']}: {error}") from error

    def add_row(self, sheet_name: str, values: dict[str, Any]) -> dict[str, Any]:
        row = self._normalize_row(sheet_name, values)
        row["_id"] = str(uuid4())
        rows = self._read_rows(sheet_name)
        rows.append(row)
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return rows[-1]

    def bulk_add_rows(self, sheet_name: str, values: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if not values:
            raise ValueError("Paste at least one record")
        rows = self._read_rows(sheet_name)
        additions = [self._normalize_row(sheet_name, item) | {"_id": str(uuid4())} for item in values]
        rows.extend(additions)
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return rows[-len(additions):]

    def preview_row_update(self, sheet_name: str, row_index: int, values: dict[str, Any]) -> dict[str, Any]:
        rows = self._read_rows(sheet_name)
        if row_index < 0 or row_index >= len(rows):
            raise IndexError("Record no longer exists")
        rows[row_index].update(self._normalize_row(sheet_name, values))
        self._calculate_formulas(sheet_name, rows)
        return rows[row_index]

    def update_row(self, sheet_name: str, row_index: int, values: dict[str, Any]) -> dict[str, Any]:
        rows = self._read_rows(sheet_name)
        if row_index < 0 or row_index >= len(rows):
            raise IndexError("Record no longer exists")
        rows[row_index].update(self._normalize_row(sheet_name, values))
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return rows[row_index]

    def bulk_update_rows(
        self,
        sheet_name: str,
        row_indices: list[int],
        values: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        if not row_indices or len(row_indices) != len(values):
            raise ValueError("Select the same number of records as pasted rows")
        if len(set(row_indices)) != len(row_indices):
            raise ValueError("A record can only be selected once")
        rows = self._read_rows(sheet_name)
        if any(index < 0 or index >= len(rows) for index in row_indices):
            raise IndexError("A selected record no longer exists")
        updates = [self._normalize_row(sheet_name, item, preserve_missing=True) for item in values]
        for index, update in zip(row_indices, updates):
            rows[index].update(update)
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return [rows[index] for index in row_indices]

    def delete_row(self, sheet_name: str, row_index: int) -> dict[str, Any]:
        rows = self._read_rows(sheet_name)
        if row_index < 0 or row_index >= len(rows):
            raise IndexError("Record no longer exists")
        record_id = rows[row_index]["_id"]
        for source_sheet, fields in self.config["sheets"].items():
            for field in fields:
                if field.get("type") == "link" and field.get("target_sheet") == sheet_name:
                    if any(row.get(field["name"]) == record_id for row in self._read_rows(source_sheet)):
                        raise ValueError("This record is linked from another list; remove the links before deleting it")
        deleted = rows.pop(row_index)
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return {"index": row_index, "row": deleted}

    def restore_row(self, sheet_name: str, row_index: int, values: dict[str, Any]) -> dict[str, Any]:
        rows = self._read_rows(sheet_name)
        insert_at = max(0, min(row_index, len(rows)))
        restored = self._normalize_row(sheet_name, values)
        restored["_id"] = values.get("_id") or str(uuid4())
        rows.insert(insert_at, restored)
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        return rows[insert_at]

    def export_workbook(self) -> bytes:
        return self.workbook_path.read_bytes()

    def list_attachments(self, sheet_name: str, record_id: str) -> list[dict[str, Any]]:
        if not any(row.get("_id") == record_id for row in self._read_rows(sheet_name)):
            raise KeyError("Record no longer exists")
        return list(self.config.setdefault("attachments", {}).get(record_id, []))

    def add_attachment(
        self,
        sheet_name: str,
        record_id: str,
        filename: str,
        content_type: str,
        content: bytes,
    ) -> dict[str, Any]:
        if not any(row.get("_id") == record_id for row in self._read_rows(sheet_name)):
            raise KeyError("Record no longer exists")
        attachment_id = str(uuid4())
        metadata = {
            "id": attachment_id,
            "name": Path(filename).name,
            "content_type": content_type or "application/octet-stream",
            "size": len(content),
        }
        if self.s3_client and self.s3_bucket:
            self.s3_client.put_object(
                Bucket=self.s3_bucket,
                Key=f"{self.s3_attachments_prefix}/{attachment_id}",
                Body=content,
                ContentType=metadata["content_type"],
            )
        else:
            (self.attachment_root / attachment_id).write_bytes(content)
        self.config.setdefault("attachments", {}).setdefault(record_id, []).append(metadata)
        self._save_config()
        return metadata

    def read_attachment(self, sheet_name: str, record_id: str, attachment_id: str) -> tuple[dict[str, Any], bytes]:
        attachments = self.list_attachments(sheet_name, record_id)
        metadata = next((item for item in attachments if item["id"] == attachment_id), None)
        if metadata is None:
            raise KeyError("Attachment not found")
        if self.s3_client and self.s3_bucket:
            response = self.s3_client.get_object(
                Bucket=self.s3_bucket,
                Key=f"{self.s3_attachments_prefix}/{attachment_id}",
            )
            content = response["Body"].read()
        else:
            content = (self.attachment_root / attachment_id).read_bytes()
        return metadata, content

    def delete_attachment(self, sheet_name: str, record_id: str, attachment_id: str) -> None:
        attachments = self.list_attachments(sheet_name, record_id)
        if not any(item["id"] == attachment_id for item in attachments):
            raise KeyError("Attachment not found")
        self._delete_attachment_blob(attachment_id)
        self.config["attachments"][record_id] = [item for item in attachments if item["id"] != attachment_id]
        self._save_config()

    def _delete_attachment_blob(self, attachment_id: str) -> None:
        if self.s3_client and self.s3_bucket:
            self.s3_client.delete_object(
                Bucket=self.s3_bucket,
                Key=f"{self.s3_attachments_prefix}/{attachment_id}",
            )
        else:
            (self.attachment_root / attachment_id).unlink(missing_ok=True)

    def _normalize_row(
        self,
        sheet_name: str,
        values: dict[str, Any],
        preserve_missing: bool = False,
    ) -> dict[str, Any]:
        fields = self._fields(sheet_name)
        row = {}
        for field in fields:
            name = field["name"]
            if field["type"] == "formula":
                continue
            if preserve_missing and name not in values:
                continue
            value = values.get(name)
            if field["type"] == "number" and value not in (None, ""):
                try:
                    value = float(value)
                except (TypeError, ValueError) as error:
                    raise ValueError(f"{name} must be a number") from error
            if field["type"] == "link" and value not in (None, ""):
                target_sheet = field["target_sheet"]
                if not any(row.get("_id") == str(value) for row in self._read_rows(target_sheet)):
                    raise ValueError(f"Choose a valid record for {name}")
                value = str(value)
            row[name] = value if value != "" else None
        return row

    def create_sheet(self, name: str, fields: list[dict[str, Any]] | None = None) -> None:
        self._validate_sheet_name(name)
        if name in self.config["sheets"]:
            raise ValueError("A list with this name already exists")
        normalized = self._normalize_fields(fields or [{"name": "Name", "type": "text"}])
        workbook = self._load_workbook()
        sheet = workbook.create_sheet(name)
        sheet.append([field["name"] for field in normalized])
        self.config["sheets"][name] = normalized
        self.config["styles"][name] = {}
        self._apply_styles(workbook, name)
        self._save_workbook(workbook)
        self._save_config()

    def add_field(self, sheet_name: str, field: dict[str, Any]) -> None:
        fields = self._fields(sheet_name)
        normalized = self._normalize_fields([field])[0]
        if any(item["name"].casefold() == normalized["name"].casefold() for item in fields):
            raise ValueError("A field with this name already exists")
        rows = self._read_rows(sheet_name)
        fields.append(normalized)
        for row in rows:
            row[normalized["name"]] = None
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        self._save_config()

    def update_field(self, sheet_name: str, field_name: str, changes: dict[str, Any]) -> None:
        fields = self._fields(sheet_name)
        field = next((item for item in fields if item["name"] == field_name), None)
        if field is None:
            raise KeyError(f"Unknown field: {field_name}")
        new_name = changes.get("name", field_name).strip()
        if not new_name:
            raise ValueError("Field name cannot be empty")
        if new_name != field_name and any(item["name"].casefold() == new_name.casefold() for item in fields):
            raise ValueError("A field with this name already exists")
        replacement = self._normalize_fields([{**field, **changes, "name": new_name}])[0]
        rows = self._read_rows(sheet_name)
        if new_name != field_name:
            for row in rows:
                row[new_name] = row.pop(field_name, None)
            style = self.config.get("styles", {}).get(sheet_name, {})
            if field_name in style:
                style[new_name] = style.pop(field_name)
        field.clear()
        field.update(replacement)
        if new_name != field_name:
            for other_fields in self.config["sheets"].values():
                for other_field in other_fields:
                    if other_field.get("type") == "link" and other_field.get("target_sheet") == sheet_name and other_field.get("target_field") == field_name:
                        other_field["target_field"] = new_name
        self._calculate_formulas(sheet_name, rows)
        self._save_rows(sheet_name, rows)
        self._save_config()

    def delete_field(self, sheet_name: str, field_name: str) -> None:
        fields = self._fields(sheet_name)
        if not any(item["name"] == field_name for item in fields):
            raise KeyError(f"Unknown field: {field_name}")
        if any(
            field.get("type") == "link" and field.get("target_sheet") == sheet_name and field.get("target_field") == field_name
            for other_fields in self.config["sheets"].values()
            for field in other_fields
        ):
            raise ValueError("This field is used as a link label by another list")
        if len(fields) == 1:
            raise ValueError("A list must have at least one field")
        rows = self._read_rows(sheet_name)
        fields[:] = [item for item in fields if item["name"] != field_name]
        self.config.get("styles", {}).get(sheet_name, {}).pop(field_name, None)
        for row in rows:
            row.pop(field_name, None)
        self._save_rows(sheet_name, rows)
        self._save_config()

    def save_styles(self, sheet_name: str, styles: dict[str, Any]) -> None:
        names = {field["name"] for field in self._fields(sheet_name)}
        if any(name not in names for name in styles):
            raise ValueError("Styles can only be applied to existing fields")
        saved_styles = self.config.setdefault("styles", {}).setdefault(sheet_name, {})
        for field_name, style in styles.items():
            saved_styles[field_name] = {**saved_styles.get(field_name, {}), **style}
        workbook = self._load_workbook()
        self._apply_styles(workbook, sheet_name)
        self._save_workbook(workbook)
        self._save_config()

    def _apply_styles(self, workbook, sheet_name: str) -> None:
        sheet = workbook[sheet_name]
        styles = self.config.get("styles", {}).get(sheet_name, {})
        for index, field in enumerate(self._fields(sheet_name), start=1):
            style = styles.get(field["name"], {})
            color = str(style.get("color", "#17211c")).lstrip("#")
            background = str(style.get("background", "#ffffff")).lstrip("#")
            border_enabled = style.get("border", False)
            border = Border(
                left=Side(style="thin", color="D9E3DD"),
                right=Side(style="thin", color="D9E3DD"),
                top=Side(style="thin", color="D9E3DD"),
                bottom=Side(style="thin", color="D9E3DD"),
            ) if border_enabled else Border()
            for row_index in range(1, sheet.max_row + 1):
                cell = sheet.cell(row_index, index)
                cell.font = Font(
                    name=cell.font.name or "Aptos",
                    size=style.get("fontSize", 11),
                    bold=style.get("bold", False),
                    italic=style.get("italic", False),
                    underline="single" if style.get("underline", False) else None,
                    color=color,
                )
                cell.fill = PatternFill(fill_type="solid", fgColor=background)
                cell.border = border

    def _normalize_fields(self, fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
        allowed_types = {"text", "number", "date", "dropdown", "formula", "link"}
        normalized = []
        seen = set()
        for field in fields:
            name = str(field.get("name", "")).strip()
            field_type = field.get("type", "text")
            if not name:
                raise ValueError("Field name cannot be empty")
            if name == ROW_ID_HEADER:
                raise ValueError(f"'{ROW_ID_HEADER}' is reserved for internal record identifiers")
            if name.casefold() in seen:
                raise ValueError("Field names must be unique")
            if field_type not in allowed_types:
                raise ValueError("Unsupported field type")
            item = {"name": name, "type": field_type}
            if field_type == "formula":
                formula = str(field.get("formula", "")).strip()
                if not formula:
                    raise ValueError("Formula fields need an expression")
                item["formula"] = formula
            if field_type == "dropdown":
                item["options"] = [str(option) for option in field.get("options", []) if str(option).strip()]
            if field_type == "link":
                target_sheet = str(field.get("target_sheet", "")).strip()
                target_field = str(field.get("target_field", "")).strip()
                if target_sheet not in self.config["sheets"]:
                    raise ValueError("Link fields must target an existing list")
                target_definition = next((target for target in self._fields(target_sheet) if target["name"] == target_field), None)
                if target_definition is None:
                    raise ValueError("Link fields must target an existing field")
                if target_definition["type"] == "link":
                    raise ValueError("Link fields cannot use another link field as their display value")
                item["target_sheet"] = target_sheet
                item["target_field"] = target_field
            normalized.append(item)
            seen.add(name.casefold())
        if not normalized:
            raise ValueError("A list must have at least one field")
        return normalized

    @staticmethod
    def _validate_sheet_name(name: str) -> None:
        if not name.strip() or len(name) > 31 or any(char in name for char in "[]:*?/\\"):
            raise ValueError("List name must be 1-31 characters and cannot contain []:*?/\\")